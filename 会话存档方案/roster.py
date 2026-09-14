"""从主管维护的 xlsx(腾讯文档导出)读取「在服务中」的学生名单。

跟旧方案基本一致,只把 chat_id 字段名改成 roomid(存档 SDK 的术语)。
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

SERVICE_WINDOW_MARGIN_DAYS = 2

IN_SERVICE_SHEETS = [
    # 换成你自己 xlsx 里的 sheet 名。只有列在这里的 sheet 会被扫描。
    "考研服务统计",
    "四六级服务统计",
]

REFUND_HINTS = ("退款", "已退")

ID_COL_NAMES = ("学生编号", "编号", "学员编号")
SUBJECT_COL_NAMES = ("科目", "服务类型")
START_COL_NAMES = ("开始时间", "开始")
EXPIRY_COL_NAMES = ("到期时间", "截至")
SINGLE_TUTOR_COL = ("接单人",)
MULTI_TUTOR_SUBJECTS = ("英语", "政治", "数学", "专业课")


@dataclass
class Roster:
    group_id: str
    raw_id: str
    subject: str
    tutor: str
    source_sheet: str
    start: date | None = None
    expiry: date | None = None
    roomid: str = ""   # 会话存档 SDK 里的群 id (wr_xxx),由 sync_rooms 回填
    raw_row: dict = field(default_factory=dict)


def _to_date(v) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y年%m月%d日", "%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


_ID_PREFIXES = ("续费年卡", "续费月卡", "续费", "续")


def _norm_id(raw: str) -> str:
    s = raw.strip()
    changed = True
    while changed:
        changed = False
        for p in _ID_PREFIXES:
            if s.startswith(p):
                s = s[len(p):].lstrip()
                changed = True
    return s


def _pick_col(cols: dict[str, int], names) -> int | None:
    for n in names:
        if n in cols:
            return cols[n]
    return None


def _extract_tutors(row: tuple, cols: dict[str, int]) -> str:
    single = _pick_col(cols, SINGLE_TUTOR_COL)
    if single is not None and row[single]:
        return str(row[single]).strip()
    parts = []
    for subj in MULTI_TUTOR_SUBJECTS:
        c = cols.get(f"{subj}接单人")
        if c is not None and row[c]:
            parts.append(f"{subj}:{str(row[c]).strip()}")
    return " / ".join(parts)


def _parse_row(row: tuple, cols: dict[str, int], today: date, sheet_name: str) -> Roster | None:
    id_col = _pick_col(cols, ID_COL_NAMES)
    if id_col is None:
        return None

    raw_id = row[id_col]
    if not raw_id:
        return None
    raw_id = str(raw_id).strip()

    if any(hint in raw_id for hint in REFUND_HINTS):
        return None

    subj_col = _pick_col(cols, SUBJECT_COL_NAMES)
    subject = ""
    if subj_col is not None and row[subj_col]:
        subject = str(row[subj_col]).strip()

    start_col = _pick_col(cols, START_COL_NAMES)
    expiry_col = _pick_col(cols, EXPIRY_COL_NAMES)
    start = _to_date(row[start_col]) if start_col is not None else None
    expiry = _to_date(row[expiry_col]) if expiry_col is not None else None

    margin = timedelta(days=SERVICE_WINDOW_MARGIN_DAYS)
    if start is not None and today < start - margin:
        return None
    if expiry is not None and today > expiry + margin:
        return None

    tutor = _extract_tutors(row, cols)

    return Roster(
        group_id=_norm_id(raw_id),
        raw_id=raw_id,
        subject=subject,
        tutor=tutor,
        source_sheet=sheet_name,
        start=start,
        expiry=expiry,
    )


def load_roster(xlsx_path: Path, today: date) -> list[Roster]:
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)

    seen: dict[str, Roster] = {}
    for sheet_name in IN_SERVICE_SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            continue
        cols = {h: i for i, h in enumerate(header) if h}

        for row in rows_iter:
            r = _parse_row(row, cols, today, sheet_name)
            if r and r.group_id and r.group_id not in seen:
                seen[r.group_id] = r

    return list(seen.values())
